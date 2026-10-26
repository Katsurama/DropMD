"""生成自检用的样例文件（packaging/samples）。开发者运行一次即可，生成结果已提交到仓库。

依赖（仅开发时）：python-docx、openpyxl、python-pptx、PySide6（用于绘制 PDF）、Pillow
"""

import os
import struct
import sys
import zlib
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
OUT = Path(__file__).resolve().parents[1] / "packaging" / "samples"
OUT.mkdir(parents=True, exist_ok=True)


def png(path: Path):
    w = h = 16
    raw = b"".join(b"\x00" + b"\x30\x80\xd0" * w for _ in range(h))

    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)

    path.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def docx():
    import docx

    img = OUT / "_tmp.png"
    png(img)
    d = docx.Document()
    d.add_heading("样例文档", 1)
    d.add_paragraph("你好，世界。こんにちは。Hello.")
    d.add_picture(str(img))
    t = d.add_table(rows=2, cols=2)
    for (r, c), v in {(0, 0): "项目", (0, 1): "数量", (1, 0): "苹果", (1, 1): "3"}.items():
        t.cell(r, c).text = v
    d.save(OUT / "sample.docx")
    img.unlink()


def xlsx():
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "表1"
    ws.append(["名称", "价格"])
    ws.append(["苹果", 5])
    ws.append(["りんご", 100])
    wb.save(OUT / "sample.xlsx")


def pptx():
    from pptx import Presentation

    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "スライドのタイトル"
    s.placeholders[1].text = "要点一\n要点二"
    p.save(OUT / "sample.pptx")


def html():
    (OUT / "sample.html").write_text(
        "<!doctype html><html><head><meta charset='utf-8'><title>样例</title></head>"
        "<body><h1>見出し</h1><p>这是一个 <b>HTML</b> 样例。</p>"
        "<table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr></table></body></html>",
        encoding="utf-8",
    )


def pdfs():
    from PySide6.QtCore import QMarginsF, QRectF, Qt
    from PySide6.QtGui import QFont, QGuiApplication, QImage, QPageSize, QPainter, QPdfWriter, QTextDocument

    app = QGuiApplication.instance() or QGuiApplication(sys.argv)  # noqa: F841

    # 文字型 PDF
    w = QPdfWriter(str(OUT / "sample_text.pdf"))
    w.setPageSize(QPageSize(QPageSize.A4))
    doc = QTextDocument()
    doc.setHtml(
        "<h1>Markdown Conversion Test</h1>"
        "<p>This PDF contains selectable text for the fast mode test.</p>"
        "<table border='1'><tr><th>Name</th><th>Value</th></tr><tr><td>alpha</td><td>1</td></tr></table>"
    )
    doc.print_(w)

    # 扫描版 PDF：整页只有一张图片，没有文字层
    img = QImage(1240, 1754, QImage.Format_RGB32)
    img.fill(Qt.white)
    p = QPainter(img)
    p.setPen(Qt.black)
    f = QFont()
    f.setPixelSize(64)
    p.setFont(f)
    p.drawText(QRectF(100, 200, 1100, 120), "DropMD OCR Test 2026")
    p.drawText(QRectF(100, 400, 1100, 120), "Scanned page sample")
    p.end()
    w2 = QPdfWriter(str(OUT / "sample_scanned.pdf"))
    w2.setPageSize(QPageSize(QPageSize.A4))
    w2.setPageMargins(QMarginsF(0, 0, 0, 0))
    w2.setResolution(150)
    p2 = QPainter(w2)
    p2.drawImage(QRectF(0, 0, w2.width(), w2.height()), img)
    p2.end()


def multilang_pdf():
    """中日英繁混合的“扫描版”PDF（整页图片），用于测试多语言 OCR。需要一款含中日文字形的字体。"""
    from PIL import Image, ImageDraw, ImageFont

    font_path = os.environ.get("CJK_FONT", "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc")
    f = ImageFont.truetype(font_path, 44)
    img = Image.new("RGB", (1240, 1754), "white")
    d = ImageDraw.Draw(img)
    for i, line in enumerate([
        "多语言识别测试：这是简体中文。",
        "これは日本語のテストです。カタカナも含む。",
        "繁體中文測試：資料處理與轉換。",
        "English OCR test, mixed languages.",
    ]):
        d.text((90, 160 + i * 110), line, fill="black", font=f)
    img.save(OUT / "sample_multilang.pdf", "PDF", resolution=150)


if __name__ == "__main__":
    docx()
    xlsx()
    pptx()
    html()
    pdfs()
    multilang_pdf()
    print("samples written to", OUT)
